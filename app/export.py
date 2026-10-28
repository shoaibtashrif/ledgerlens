"""Ledger export (Excel / CSV) and summary metrics."""
from __future__ import annotations

import csv
import io
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

LEDGER_STATUSES = ["auto_approved", "approved"]
COLUMNS = ["invoice_date", "vendor_name", "invoice_number", "category", "currency",
           "subtotal", "tax", "total", "due_date", "status", "confidence", "provider", "filename"]


def ledger_rows(docs: list[dict]) -> list[dict]:
    rows = []
    for d in docs:
        data = d.get("data") or {}
        rows.append({
            "invoice_date": data.get("invoice_date"), "vendor_name": data.get("vendor_name"),
            "invoice_number": data.get("invoice_number"), "category": d.get("category"),
            "currency": data.get("currency"), "subtotal": data.get("subtotal"), "tax": data.get("tax"),
            "total": data.get("total"), "due_date": data.get("due_date"), "status": d["status"],
            "confidence": d.get("confidence"), "provider": d.get("provider"), "filename": d["filename"],
        })
    return sorted(rows, key=lambda r: r["invoice_date"] or "")


def to_csv(docs: list[dict]) -> str:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=COLUMNS)
    w.writeheader()
    w.writerows(ledger_rows(docs))
    return buf.getvalue()


def to_xlsx(ledger: list[dict], exceptions: list[dict]) -> bytes:
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="4F46E5")

    def sheet(ws, header, rows):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(40, max(10, *(len(str(c.value or "")) + 2 for c in col)))

    ws = wb.active
    ws.title = "Ledger"
    sheet(ws, COLUMNS, [[r[c] for c in COLUMNS] for r in ledger_rows(ledger)])

    items = []
    for d in ledger:
        data = d.get("data") or {}
        for li in data.get("line_items") or []:
            items.append([data.get("vendor_name"), data.get("invoice_number"), li.get("description"),
                          li.get("quantity"), li.get("unit_price"), li.get("amount")])
    sheet(wb.create_sheet("Line items"), ["vendor", "invoice_number", "description", "qty", "unit_price", "amount"], items)

    summ = summary(ledger)
    rows = [[cat, cur, round(v, 2)] for (cat, cur), v in sorted(summ["by_category_currency"].items())]
    sheet(wb.create_sheet("Summary"), ["category", "currency", "total"], rows)

    exc = [[d["filename"], d["status"], "; ".join(i["message"] for i in d.get("issues") or [])] for d in exceptions]
    sheet(wb.create_sheet("Exceptions"), ["file", "status", "issues"], exc)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def summary(ledger: list[dict]) -> dict:
    by_cat_cur: dict = defaultdict(float)
    by_vendor: dict = defaultdict(float)
    by_cur: dict = defaultdict(float)
    by_month: dict = defaultdict(float)
    for d in ledger:
        data = d.get("data") or {}
        total = data.get("total") or 0
        cur = data.get("currency") or "?"
        by_cat_cur[(d.get("category") or "Uncategorized", cur)] += total
        by_vendor[f"{data.get('vendor_name') or 'Unknown'} ({cur})"] += total
        by_cur[cur] += total
        by_month[((data.get("invoice_date") or "unknown")[:7], cur)] += total
    return {"by_category_currency": by_cat_cur, "by_vendor": by_vendor, "by_currency": by_cur, "by_month": by_month}


def metrics(all_docs: list[dict]) -> dict:
    n = len(all_docs)
    counts: dict = defaultdict(int)
    for d in all_docs:
        counts[d["status"]] += 1
    processed = [d for d in all_docs if d["status"] != "failed"]
    reviewed = [d for d in all_docs if d.get("reviewed_at")]
    touched = [d for d in reviewed if d.get("edited_fields")]
    lat = [d["latency_ms"] for d in all_docs if d.get("latency_ms")]
    costs = [d["cost_usd"] for d in all_docs if d.get("cost_usd") is not None]
    providers: dict = defaultdict(int)
    for d in processed:
        providers[d.get("provider") or "?"] += 1
    return {
        "total": n, "status_counts": dict(counts),
        "straight_through_rate": round(counts["auto_approved"] / n, 3) if n else 0,
        "manual_touch_rate": round(len(touched) / n, 3) if n else 0,
        "reviewed": len(reviewed), "edited": len(touched),
        "avg_latency_s": round(sum(lat) / len(lat) / 1000, 2) if lat else 0,
        "tokens": sum((d.get("tokens_in") or 0) + (d.get("tokens_out") or 0) for d in all_docs),
        "est_cost_usd": round(sum(costs), 4) if costs else None,
        "providers": dict(providers),
    }
